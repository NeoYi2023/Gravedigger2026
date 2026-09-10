# -*- coding: utf-8 -*-
"""Insert OffScreenSpawnHint* rows into CombatConstantConfig Excel, then Bake CSV."""
from __future__ import annotations

import io
import math
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "Gravedigger2026" / "Assets" / "ConfigTables"
EXCEL_NAME = "通用_常量表_Combat_CombatConstantConfig.xlsx"
CSV_NAME = "Combat_CombatConstantConfig.csv"
ANCHOR_KEYS = (
    "SearchExtractDecisionAutoLeaveSeconds",
    "PushMapCameraIntroWaypointDwellSeconds",
)

EN_COL = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")
MAX_SCAN = 3
_MAX_DECIMAL_PLACES = 10

ROWS = [
    (
        "OffScreenSpawnHintDurationSeconds",
        "离屏来袭提示时长",
        1.5,
        "UI-034: edge icon blink total seconds then hide",
        "UI-034：边缘图标闪烁总秒数后消失",
    ),
    (
        "OffScreenSpawnHintBlinkPeriodSeconds",
        "离屏来袭闪烁周期",
        0.25,
        "UI-034: alpha on/off period seconds",
        "UI-034：alpha 一亮一灭周期秒",
    ),
    (
        "OffScreenSpawnHintEdgeMarginPx",
        "离屏来袭边距像素",
        48,
        "UI-034: icon center margin from screen edge (ref 1920x1080 px)",
        "UI-034：图标中心距参考分辨率屏幕边的像素边距",
    ),
    (
        "OffScreenSpawnHintIconSizePx",
        "离屏来袭图标边长",
        72,
        "UI-034: EnemyAttack_1 sizeDelta edge (ref 1920x1080)",
        "UI-034：EnemyAttack_1 sizeDelta 边长（参考 1920×1080）",
    ),
]


def format_numeric_for_csv(number: float) -> str:
    if not math.isfinite(number):
        return str(number)
    if abs(number - round(number)) < 1e-9 and abs(number) < 1e15:
        return str(int(round(number)))
    quant = Decimal(1).scaleb(-_MAX_DECIMAL_PLACES)
    rounded_dec = Decimal(repr(number)).quantize(quant, rounding=ROUND_HALF_UP)
    rounded = float(rounded_dec)
    if abs(rounded - round(rounded)) < 1e-12 and abs(rounded) < 1e15:
        return str(int(round(rounded)))
    text = f"{rounded:.{_MAX_DECIMAL_PLACES}f}".rstrip("0").rstrip(".")
    return text if text not in ("", "-") else "0"


def sanitize_embedded_float_noise(text: str) -> str:
    if not text or "." not in text:
        return text

    def repl(match: re.Match[str]) -> str:
        token = match.group(0)
        frac = token.split(".", 1)[1]
        looks_noisy = len(frac) > 10 or "000000" in frac or "999999" in frac
        if not looks_noisy:
            return token
        try:
            number = float(token)
        except ValueError:
            return token
        return format_numeric_for_csv(number)

    return re.sub(r"-?\d+\.\d+", repl, text)


def cell_to_csv_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return format_numeric_for_csv(value)
    if isinstance(value, Decimal):
        return format_numeric_for_csv(float(value))
    return sanitize_embedded_float_noise(str(value))


def is_english_header(row: list[str]) -> bool:
    saw = False
    for cell in row:
        if cell is None or str(cell).strip() == "":
            continue
        saw = True
        if not EN_COL.match(str(cell).strip()):
            return False
    return saw


def find_header_index(rows: list[list[str]]) -> int:
    for i, row in enumerate(rows[:MAX_SCAN]):
        if is_english_header(row):
            return i
    raise ValueError("no English header in first 3 rows")


def escape_csv_field(value: str) -> str:
    if any(ch in value for ch in [",", '"', "\n", "\r"]):
        return '"' + value.replace('"', '""') + '"'
    return value


def build_csv(rows: list[list[str]]) -> str:
    lines: list[str] = []
    for r_i, row in enumerate(rows):
        if r_i > 0 and all(not (c or "").strip() for c in row):
            continue
        lines.append(",".join(escape_csv_field(c or "") for c in row))
    return "\n".join(lines) + ("\n" if lines else "")


def insert_rows(xlsx_path: Path) -> None:
    wb = load_workbook(xlsx_path)
    ws = wb.active
    header3 = [ws.cell(3, c).value for c in range(1, 6)]
    if header3 != ["ConstantKey", "ConstantKeyZh", "Value", "Comment", "CommentZh"]:
        raise SystemExit(f"unexpected header row 3 in {xlsx_path}: {header3}")

    existing = {ws.cell(r, 1).value for r in range(1, ws.max_row + 1)}
    if all(k[0] in existing for k in ROWS):
        print(f"SKIP already has OffScreenSpawnHint keys: {xlsx_path}")
        wb.close()
        return

    insert_at = None
    used_anchor = None
    for anchor in ANCHOR_KEYS:
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == anchor:
                insert_at = r + 1
                used_anchor = anchor
                break
        if insert_at is not None:
            break
    if insert_at is None:
        raise SystemExit(f"anchors {ANCHOR_KEYS} not found in {xlsx_path}")
    print(f"anchor={used_anchor} insert_at={insert_at} for {xlsx_path.name}")

    n = len(ROWS)
    ws.insert_rows(insert_at, amount=n)
    for i, row in enumerate(ROWS):
        for c, val in enumerate(row, start=1):
            ws.cell(insert_at + i, c).value = val

    wb.save(xlsx_path)
    wb.close()
    print(f"INSERTED {n} rows at Excel row {insert_at}: {xlsx_path}")


def bake(excel_dir: Path, csv_dir: Path) -> None:
    xlsx = excel_dir / EXCEL_NAME
    with open(xlsx, "rb") as raw:
        data = raw.read()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        cells = [cell_to_csv_text(v) for v in row]
        rows.append(cells)
    wb.close()
    header_i = find_header_index(rows)
    if header_i < 2:
        raise SystemExit(f"expected 3-row header, got English at row {header_i}")
    text = build_csv(rows[header_i:])
    out = csv_dir / CSV_NAME
    csv_dir.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8", newline="\n")
    keys = [ln.split(",", 1)[0] for ln in text.splitlines()[1:]]
    missing = [k[0] for k in ROWS if k[0] not in keys]
    if missing:
        raise SystemExit(f"Bake missing keys {missing} in {out}")
    print(f"BAKED {xlsx.name} → {out}")


def main() -> int:
    targets = [
        (ROOT / "Excel" / EXCEL_NAME, ROOT / "Excel", ROOT / "Csv"),
        (ROOT / "Mode2" / "Excel" / EXCEL_NAME, ROOT / "Mode2" / "Excel", ROOT / "Mode2" / "Csv"),
    ]
    for xlsx, excel_dir, csv_dir in targets:
        if not xlsx.is_file():
            raise SystemExit(f"missing {xlsx}")
        insert_rows(xlsx)
        bake(excel_dir, csv_dir)
    return 0


if __name__ == "__main__":
    sys.exit(main())
