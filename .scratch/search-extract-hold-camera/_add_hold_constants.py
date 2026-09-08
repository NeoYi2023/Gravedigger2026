# -*- coding: utf-8 -*-
"""Insert SearchExtractHold* rows into CombatConstantConfig Excel, then Bake CSV (SPEC_04 §14.7)."""
from __future__ import annotations

import io
import math
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "Gravedigger2026" / "Assets" / "ConfigTables"
EXCEL_NAME = "通用_常量表_Combat_CombatConstantConfig.xlsx"
CSV_NAME = "Combat_CombatConstantConfig.csv"
ANCHOR_KEY = "PushMapCameraIntroWaypointDwellSeconds"

EN_COL = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")
MAX_SCAN = 3
_MAX_DECIMAL_PLACES = 10

HOLD_ROWS = [
    (
        "SearchExtractHoldOrthoSizeMin",
        "搜打撤守点Size下限",
        3,
        "HoldFraming ortho Size floor (intersect CameraOrthoSizeMin)",
        "HoldFraming 正交 Size 下限（再与 CameraOrthoSizeMin 取交）",
    ),
    (
        "SearchExtractHoldOrthoSizeMax",
        "搜打撤守点Size上限",
        8,
        "HoldFraming ortho Size ceiling (intersect CameraOrthoSizeMax)",
        "HoldFraming 正交 Size 上限（再与 CameraOrthoSizeMax 取交）",
    ),
    (
        "SearchExtractHoldMaxPanRadius",
        "守点中心偏移半径",
        8,
        "Max look-at distance from current Objective world XZ",
        "look-at 距当前 Objective 世界 XZ 最大半径；圈外忠诚兵允许出镜",
    ),
    (
        "SearchExtractHoldViewportPad",
        "守点视口边垫",
        0.08,
        "Outer viewport pad on four sides (0-0.5)",
        "外框四边视口垫（0～0.5）",
    ),
    (
        "SearchExtractHoldTopHudPad",
        "守点顶HUD垫",
        0.12,
        "Extra top pad for CountdownBar / CombatIndicator",
        "外框顶边额外垫（CountdownBar / CombatIndicator）",
    ),
    (
        "SearchExtractHoldInnerPad",
        "守点内框垫",
        0.18,
        "Inner pad; tighten only after all stay inside for dwell",
        "内框边垫；全体在内框内满滞留才允许收紧",
    ),
    (
        "SearchExtractHoldZoomInDelaySeconds",
        "守点拉近滞留秒",
        0.8,
        "Dwell seconds before allowing Size shrink / recenter",
        "内框滞留后才允许减小 Size / 回中",
    ),
    (
        "SearchExtractHoldSmoothTimeOut",
        "守点拉远平滑秒",
        0.2,
        "SmoothDamp when expanding Size / pan-out",
        "Size/中心拉远（扩大）SmoothDamp",
    ),
    (
        "SearchExtractHoldSmoothTimeIn",
        "守点拉近平滑秒",
        0.55,
        "SmoothDamp when tightening; must be > Out",
        "Size/中心拉近（收紧）SmoothDamp；须大于 Out",
    ),
    (
        "SearchExtractHoldSampleInterval",
        "守点采样间隔秒",
        0.2,
        "Low-frequency viewport AABB sample interval seconds",
        "视口 AABB 低频采样间隔",
    ),
    (
        "SearchExtractHoldObjectiveBias",
        "守点中心偏向",
        0.65,
        "look-at = lerp(soldier center, Objective, bias); 1=pin Objective",
        "look-at = lerp(士兵中心, Objective, bias)；1=钉 Objective",
    ),
]


def format_numeric_for_csv(number: float) -> str:
    if not math.isfinite(number):
        return str(number)
    if abs(number - round(number)) < 1e-9 and abs(number) < 1e15:
        return str(int(round(number)))
    quant = Decimal(1).scaleb(-_MAX_DECIMAL_PLACES)
    rounded_dec = Decimal(repr(number)).quantize(quant, rounding=ROUND_HALF_UP)
    rounded = float(rounded_dec)
    if abs(rounded - round(rounded)) < 1e-12 and abs(rounded) < 1e15:
        return str(int(round(rounded)))
    text = f"{rounded:.{_MAX_DECIMAL_PLACES}f}".rstrip("0").rstrip(".")
    return text if text not in ("", "-") else "0"


def sanitize_embedded_float_noise(text: str) -> str:
    if not text or "." not in text:
        return text

    def repl(match: re.Match[str]) -> str:
        token = match.group(0)
        frac = token.split(".", 1)[1]
        looks_noisy = len(frac) > 10 or "000000" in frac or "999999" in frac
        if not looks_noisy:
            return token
        try:
            number = float(token)
        except ValueError:
            return token
        return format_numeric_for_csv(number)

    return re.sub(r"-?\d+\.\d+", repl, text)


def cell_to_csv_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return format_numeric_for_csv(value)
    if isinstance(value, Decimal):
        return format_numeric_for_csv(float(value))
    return sanitize_embedded_float_noise(str(value))


def is_english_header(row: list[str]) -> bool:
    saw = False
    for cell in row:
        if cell is None or str(cell).strip() == "":
            continue
        saw = True
        if not EN_COL.match(str(cell).strip()):
            return False
    return saw


def find_header_index(rows: list[list[str]]) -> int:
    for i, row in enumerate(rows[:MAX_SCAN]):
        if is_english_header(row):
            return i
    raise ValueError("no English header in first 3 rows")


def escape_csv_field(value: str) -> str:
    if any(ch in value for ch in [",", '"', "\n", "\r"]):
        return '"' + value.replace('"', '""') + '"'
    return value


def build_csv(rows: list[list[str]]) -> str:
    lines: list[str] = []
    for r_i, row in enumerate(rows):
        if r_i > 0 and all(not (c or "").strip() for c in row):
            continue
        lines.append(",".join(escape_csv_field(c or "") for c in row))
    return "\n".join(lines) + ("\n" if lines else "")


def insert_hold_rows(xlsx_path: Path) -> None:
    wb = load_workbook(p := xlsx_path)
    ws = wb.active
    header1 = [ws.cell(1, c).value for c in range(1, 6)]
    header2 = [ws.cell(2, c).value for c in range(1, 6)]
    header3 = [ws.cell(3, c).value for c in range(1, 6)]
    if header3 != ["ConstantKey", "ConstantKeyZh", "Value", "Comment", "CommentZh"]:
        raise SystemExit(f"unexpected header row 3 in {p}: {header3}")
    if not header1[0] or not header2[0]:
        raise SystemExit(f"missing doc rows 1-2 in {p}")

    existing = {ws.cell(r, 1).value for r in range(1, ws.max_row + 1)}
    if all(k[0] in existing for k in HOLD_ROWS):
        print(f"SKIP already has Hold keys: {xlsx_path}")
        wb.close()
        return

    insert_at = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == ANCHOR_KEY:
            insert_at = r + 1
            break
    if insert_at is None:
        raise SystemExit(f"anchor {ANCHOR_KEY} not found in {p}")

    n = len(HOLD_ROWS)
    ws.insert_rows(insert_at, amount=n)
    for i, row in enumerate(HOLD_ROWS):
        for c, val in enumerate(row, start=1):
            ws.cell(insert_at + i, c).value = val

    # Doc rows must still be present.
    if [ws.cell(1, c).value for c in range(1, 6)] != header1:
        raise SystemExit(f"row 1 mutated in {p}")
    if [ws.cell(2, c).value for c in range(1, 6)] != header2:
        raise SystemExit(f"row 2 mutated in {p}")

    wb.save(xlsx_path)
    wb.close()
    print(f"INSERTED {n} Hold rows at Excel row {insert_at}: {xlsx_path}")


def bake(excel_dir: Path, csv_dir: Path) -> None:
    xlsx = excel_dir / EXCEL_NAME
    with open(xlsx, "rb") as raw:
        data = raw.read()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        cells = [cell_to_csv_text(v) for v in row]
        rows.append(cells)
    wb.close()
    if not rows:
        raise SystemExit(f"empty sheet {xlsx}")
    max_cols = max(len(r) for r in rows)
    for r in rows:
        if len(r) < max_cols:
            r.extend([""] * (max_cols - len(r)))
    header_i = find_header_index(rows)
    if header_i < 2:
        raise SystemExit(f"expected 3-row header, got English at row {header_i} in {xlsx}")
    text = build_csv(rows[header_i:])
    out = csv_dir / CSV_NAME
    csv_dir.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8", newline="\n")
    keys = [ln.split(",", 1)[0] for ln in text.splitlines()[1:]]
    missing = [k[0] for k in HOLD_ROWS if k[0] not in keys]
    if missing:
        raise SystemExit(f"Bake missing keys {missing} in {out}")
    print(f"BAKED {xlsx.name} → {out} (doc_rows={header_i}, hold_keys={len(HOLD_ROWS)})")


def main() -> int:
    targets = [
        (ROOT / "Excel" / EXCEL_NAME, ROOT / "Excel", ROOT / "Csv"),
        (ROOT / "Mode2" / "Excel" / EXCEL_NAME, ROOT / "Mode2" / "Excel", ROOT / "Mode2" / "Csv"),
    ]
    for xlsx, excel_dir, csv_dir in targets:
        if not xlsx.is_file():
            raise SystemExit(f"missing {xlsx}")
        insert_hold_rows(xlsx)
        bake(excel_dir, csv_dir)
    return 0


if __name__ == "__main__":
    sys.exit(main())
