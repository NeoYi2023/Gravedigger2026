# -*- coding: utf-8 -*-
"""D-100 slice 02: Mode2 Class_BaseWarrior TargetTypeScores adds 可破坏障碍物 (SPEC_04 §14.7)."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "Gravedigger2026" / "Assets" / "ConfigTables"
CLASS_XLSX = "制造_职业配置表_Manufacture_ClassConfig.xlsx"
WARRIOR_ID = "Class_BaseWarrior"
NEW_SCORES = "敌方单位;1|可破坏障碍物;1"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bake_config_tables_py import (  # noqa: E402
    build_csv,
    excel_to_csv_base,
    find_header_index,
    read_sheet_rows,
)


def find_en_header_row(ws) -> int:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True), 1):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if "ClassId" in cells:
            return i
    raise RuntimeError("no EN header in first 3 rows")


def header_list(ws, header_row: int) -> list[str]:
    row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    names = [str(c).strip() if c is not None else "" for c in row]
    while names and names[-1] == "":
        names.pop()
    return names


def col_index(names: list[str], en: str) -> int:
    return names.index(en) + 1


def patch_mode2_warrior(path: Path) -> None:
    wb = load_workbook(path)
    ws = wb.active
    header_row = find_en_header_row(ws)
    names = header_list(ws, header_row)
    id_col = col_index(names, "ClassId")
    score_col = col_index(names, "TargetTypeScores")
    patched = False
    for r in range(header_row + 1, ws.max_row + 1):
        class_id = ws.cell(r, id_col).value
        if class_id is None or str(class_id).strip() != WARRIOR_ID:
            continue
        ws.cell(r, score_col, NEW_SCORES)
        patched = True
        break
    if not patched:
        raise RuntimeError(f"{path}: missing {WARRIOR_ID}")
    wb.save(path)
    print(f"patched {path}")


def bake_one(excel_dir: Path, csv_dir: Path, xlsx_name: str) -> None:
    xlsx = excel_dir / xlsx_name
    rows = read_sheet_rows(xlsx)
    header_i = find_header_index(rows)
    csv_name = excel_to_csv_base(xlsx.stem) + ".csv"
    dest = csv_dir / csv_name
    dest.write_text(build_csv(rows, header_i), encoding="utf-8-sig")
    print(f"baked {dest}")


def main() -> None:
    mode2_xlsx = ROOT / "Mode2" / "Excel" / CLASS_XLSX
    patch_mode2_warrior(mode2_xlsx)
    bake_one(ROOT / "Mode2" / "Excel", ROOT / "Mode2" / "Csv", CLASS_XLSX)


if __name__ == "__main__":
    main()
