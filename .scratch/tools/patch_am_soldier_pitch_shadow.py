# -*- coding: utf-8 -*-
"""Patch AutoMfgSoldierPitchFactor + ShadowAlpha=0 on Mode1/Mode2 CombatConstant Excel."""
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "Gravedigger2026" / "Assets" / "ConfigTables"
PATHS = [
    ROOT / "Excel" / "通用_常量表_Combat_CombatConstantConfig.xlsx",
    ROOT / "Mode2" / "Excel" / "通用_常量表_Combat_CombatConstantConfig.xlsx",
]

PITCH_ROW = (
    "AutoMfgSoldierPitchFactor",
    "自动制造士兵行间距系数",
    0.5,
    "StepB row pitch = MaxEdge x VisualScale x Factor",
    "StepB 行间距=MaxEdge×VisualScale×本值（样例0.5→256）",
)


def patch(path: Path) -> None:
    wb = load_workbook(path)
    ws = wb.active
    alpha_row = None
    visual_row = None
    pitch_exists = False
    for i, row in enumerate(ws.iter_rows(min_row=1, max_col=5), 1):
        key = row[0].value
        if key == "AutoMfgSoldierShadowAlpha":
            alpha_row = i
            row[2].value = 0
            row[3].value = "Foot-shadow black Alpha 0-1; 0=Demo hide"
            row[4].value = "脚下影子黑色 Alpha（0～1；0=Demo 隐藏）"
        elif key == "AutoMfgSoldierVisualScale":
            visual_row = i
        elif key == "AutoMfgSoldierPitchFactor":
            pitch_exists = True
            row[2].value = 0.5
    if not pitch_exists:
        insert_at = (visual_row + 1) if visual_row else ((alpha_row + 1) if alpha_row else ws.max_row + 1)
        ws.insert_rows(insert_at)
        for col, val in enumerate(PITCH_ROW, 1):
            ws.cell(insert_at, col, val)
    wb.save(path)
    print("patched", path)


def main() -> None:
    for path in PATHS:
        patch(path)
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if row and row[0] in (
                "AutoMfgSoldierShadowAlpha",
                "AutoMfgSoldierPitchFactor",
                "AutoMfgSoldierVisualScale",
            ):
                print(" ", i, row)


if __name__ == "__main__":
    main()
