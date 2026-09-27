# -*- coding: utf-8 -*-
"""TFG-01: add FormationLevel composite key and wedge level-5 row. Bake Mode1+Mode2 CSV."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tf01_tactical_formation_config import (  # noqa: E402
    FORMATION_CSV,
    FORMATION_XLSX,
    ROOT,
    bake,
    header_map,
)

LEVEL_STAT = "Stat=Strength|Mul=1.30"
WEDGE_ID = "Form_Wedge_01"


def patch(xlsx: Path) -> None:
    wb = load_workbook(xlsx)
    ws = wb.active
    cols = header_map(ws)
    if "FormationLevel" not in cols:
        insert_at = cols["FormationId"] + 1
        ws.insert_cols(insert_at)
        ws.cell(1, insert_at, "阵型等级")
        ws.cell(2, insert_at, "复合主键之一；≥1；同 Id 用它区分属性")
        ws.cell(3, insert_at, "FormationLevel")
        print(f"  inserted FormationLevel at col {insert_at}")
        cols = header_map(ws)

    id_col = cols["FormationId"]
    level_col = cols["FormationLevel"]
    stat_col = cols["StatModifiers"]
    wedge_level1 = None
    has_level5 = False
    for r in range(4, ws.max_row + 1):
        raw_id = ws.cell(r, id_col).value
        if raw_id is None or str(raw_id).strip() == "":
            continue
        formation_id = str(raw_id).strip()
        level = ws.cell(r, level_col).value
        if level is None or str(level).strip() == "":
            ws.cell(r, level_col, 1)
            level = 1
            print(f"  {xlsx.name} row {r} {formation_id} FormationLevel=1")
        level = int(level)
        if formation_id == WEDGE_ID and level == 1:
            wedge_level1 = r
        if formation_id == WEDGE_ID and level == 5:
            has_level5 = True

    if wedge_level1 is None:
        raise SystemExit(f"{xlsx}: missing {WEDGE_ID} level 1")
    if not has_level5:
        new_r = ws.max_row + 1
        for c in range(1, ws.max_column + 1):
            ws.cell(new_r, c, ws.cell(wedge_level1, c).value)
        ws.cell(new_r, level_col, 5)
        ws.cell(new_r, stat_col, LEVEL_STAT)
        print(f"  appended {WEDGE_ID} level 5 at row {new_r}")
    wb.save(xlsx)
    print(f"Patched {xlsx}")


def main() -> None:
    pairs = [
        (ROOT / "Excel", ROOT / "Csv"),
        (ROOT / "Mode2" / "Excel", ROOT / "Mode2" / "Csv"),
    ]
    for excel_dir, csv_dir in pairs:
        xlsx = excel_dir / FORMATION_XLSX
        patch(xlsx)
        bake(xlsx, csv_dir / FORMATION_CSV)


if __name__ == "__main__":
    main()
