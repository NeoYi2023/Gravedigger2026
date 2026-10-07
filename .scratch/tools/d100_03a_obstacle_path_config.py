# -*- coding: utf-8 -*-
"""D-100 slice 03a: MonsterConfig.ObstaclePathMode + COC spawn sample (SPEC_04 §14.7)."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "Gravedigger2026" / "Assets" / "ConfigTables"
MONSTER_XLSX = "防守_怪物配置表_Defend_MonsterConfig.xlsx"
SPAWN_XLSX = "COC_刷怪配置表_Coc_CocSpawnConfig.xlsx"

MONSTER_COL = (
    "ObstaclePathMode",
    "障碍寻路模式",
    "仅COC：视为障碍|视为目标|无视障碍；缺列/空→视为障碍；非法失败",
)

MODE2_PATH_BY_ID = {
    "Monster_01": "视为障碍",
    "Monster_02": "无视障碍",
    "Monster_03": "视为目标",
}

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bake_config_tables_py import (  # noqa: E402
    build_csv,
    excel_to_csv_base,
    find_header_index,
    read_sheet_rows,
)


def find_en_header_row(ws) -> int:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True), 1):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if any(c in ("MonsterId", "GameplayConfigId") for c in cells):
            return i
    raise RuntimeError("no EN header in first 3 rows")


def header_list(ws, header_row: int) -> list[str]:
    row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    names = [str(c).strip() if c is not None else "" for c in row]
    while names and names[-1] == "":
        names.pop()
    return names


def append_column(ws, header_row: int, en: str, zh: str, note: str) -> None:
    names = header_list(ws, header_row)
    if en in names:
        return
    insert_at = len(names) + 1
    ws.insert_cols(insert_at)
    if header_row >= 3:
        ws.cell(1, insert_at, zh)
        ws.cell(2, insert_at, note)
    ws.cell(header_row, insert_at, en)


def col_index(names: list[str], en: str) -> int:
    return names.index(en) + 1


def patch_monster(path: Path, fill_demo: bool) -> None:
    wb = load_workbook(path)
    ws = wb.active
    header_row = find_en_header_row(ws)
    append_column(ws, header_row, *MONSTER_COL)
    names = header_list(ws, header_row)
    id_col = col_index(names, "MonsterId")
    mode_col = col_index(names, "ObstaclePathMode")
    for r in range(header_row + 1, ws.max_row + 1):
        monster_id = ws.cell(r, id_col).value
        if monster_id is None or not str(monster_id).strip():
            continue
        if not fill_demo:
            ws.cell(r, mode_col, "")
            continue
        key = str(monster_id).strip()
        ws.cell(r, mode_col, MODE2_PATH_BY_ID.get(key, "视为障碍"))
    wb.save(path)
    print(f"OK monster fill_demo={fill_demo} {path}")


def patch_spawn(path: Path) -> None:
    wb = load_workbook(path)
    ws = wb.active
    header_row = find_en_header_row(ws)
    names = header_list(ws, header_row)
    required = (
        "GameplayConfigId",
        "SpawnPointId",
        "MonsterId",
        "SpawnCount",
        "SpawnRole",
        "SpawnOrder",
    )
    cols = {en: col_index(names, en) for en in required}
    for r in range(ws.max_row, header_row, -1):
        ws.delete_rows(r)

    rows = []
    for gameplay in ("Coc_Normal", "Coc_Hard", "Coc_Hell"):
        rows.append((gameplay, "SP_01", "Monster_01", 1, "Normal", 1))
        rows.append((gameplay, "SP_01", "Monster_02", 1, "Normal", 2))
        rows.append((gameplay, "SP_01", "Monster_03", 1, "Normal", 3))
        rows.append((gameplay, "SP_Mini", "Monster_01", 1, "MiniBoss", 1))
        rows.append((gameplay, "SP_Final", "Monster_01", 1, "FinalBoss", 1))

    for i, row in enumerate(rows):
        excel_row = header_row + 1 + i
        ws.cell(excel_row, cols["GameplayConfigId"], row[0])
        ws.cell(excel_row, cols["SpawnPointId"], row[1])
        ws.cell(excel_row, cols["MonsterId"], row[2])
        ws.cell(excel_row, cols["SpawnCount"], row[3])
        ws.cell(excel_row, cols["SpawnRole"], row[4])
        ws.cell(excel_row, cols["SpawnOrder"], row[5])

    wb.save(path)
    print(f"OK spawn {path} rows={len(rows)}")


def bake_one(excel_dir: Path, csv_dir: Path, xlsx_name: str) -> None:
    xlsx = excel_dir / xlsx_name
    rows = read_sheet_rows(xlsx)
    header_i = find_header_index(rows)
    text = build_csv(rows[header_i:])
    out = csv_dir / f"{excel_to_csv_base(xlsx.stem)}.csv"
    out.write_text(text, encoding="utf-8", newline="\n")
    header = text.splitlines()[0] if text else ""
    print(f"BAKE {xlsx.name} → {out.name} header={header}")


def main() -> None:
    mode1_excel = ROOT / "Excel"
    mode1_csv = ROOT / "Csv"
    mode2_excel = ROOT / "Mode2" / "Excel"
    mode2_csv = ROOT / "Mode2" / "Csv"

    patch_monster(mode1_excel / MONSTER_XLSX, fill_demo=False)
    patch_monster(mode2_excel / MONSTER_XLSX, fill_demo=True)
    patch_spawn(mode2_excel / SPAWN_XLSX)

    bake_one(mode1_excel, mode1_csv, MONSTER_XLSX)
    bake_one(mode2_excel, mode2_csv, MONSTER_XLSX)
    bake_one(mode2_excel, mode2_csv, SPAWN_XLSX)


if __name__ == "__main__":
    main()
