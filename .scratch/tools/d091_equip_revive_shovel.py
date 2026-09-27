# -*- coding: utf-8 -*-
"""D-091: Equip_ReviveShovel L1–5 on Mode2 ProtagonistEquipment Excel + Bake."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "Gravedigger2026" / "Assets" / "ConfigTables"
sys.path.insert(0, str(Path(__file__).resolve().parent))
from bake_config_tables_py import (  # noqa: E402
    build_csv,
    excel_to_csv_base,
    find_header_index,
    read_sheet_rows,
)

EQUIP_XLSX = "主角_装备配置表_Protagonist_ProtagonistEquipmentConfig.xlsx"
EQUIP_ID = "Equip_ReviveShovel"

CHANCES = {1: "0.2", 2: "0.4", 3: "0.6", 4: "0.8", 5: "1"}
CONVERT = {1: "1", 2: "2", 3: "3", 4: "4", 5: "5"}
PCT = {1: "20", 2: "40", 3: "60", 4: "80", 5: "100"}


def find_en_header(ws) -> tuple[int, dict[str, int]]:
    for r in range(1, min(4, ws.max_row + 1)):
        cells = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        texts = [str(v).strip() if v is not None else "" for v in cells]
        if "EquipId" in texts:
            cols = {t: i + 1 for i, t in enumerate(texts) if t}
            return r, cols
    raise SystemExit("no EN header EquipId in first 3 rows")


def patch_equip_xlsx(path: Path) -> None:
    wb = load_workbook(path)
    ws = wb.active
    header_row, cols = find_en_header(ws)
    id_col = cols["EquipId"]
    existing: dict[int, int] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        if ws.cell(r, id_col).value != EQUIP_ID:
            continue
        lvl_raw = ws.cell(r, cols["EquipLevel"]).value
        if lvl_raw is None or str(lvl_raw).strip() == "":
            continue
        existing[int(lvl_raw)] = r

    def row_values(lvl: int) -> list[object]:
        effect = f"DigOnGraveClear_{CHANCES[lvl]}|DigLightningPreviewSec_2"
        desc = f"挖开坟墓时有{PCT[lvl]}%概率直接产出士兵"
        return [
            EQUIP_ID,
            lvl,
            "复活铲",
            "Icon_IronShovel",
            None if lvl == 5 else 1,
            int(CONVERT[lvl]),
            "Dig",
            effect,
            desc,
        ]

    if existing:
        for lvl, row_i in existing.items():
            values = row_values(lvl)
            for c, value in enumerate(values, start=1):
                ws.cell(row_i, c).value = value
        print(f"PATCHED {len(existing)} {EQUIP_ID} rows: {path}")
    else:
        for lvl in range(1, 6):
            ws.append(row_values(lvl))
        print(f"APPENDED 5 {EQUIP_ID} rows: {path}")

    wb.save(path)


def bake_one(excel_dir: Path, csv_dir: Path, xlsx_name: str) -> None:
    xlsx = excel_dir / xlsx_name
    rows = read_sheet_rows(xlsx)
    header_i = find_header_index(rows)
    export = rows[header_i:]
    text = build_csv(export)
    out = csv_dir / f"{excel_to_csv_base(xlsx.stem)}.csv"
    out.write_text(text, encoding="utf-8", newline="\n")
    print(f"BAKE {xlsx.name} → {out.name}")


def main() -> None:
    mode2_equip = ROOT / "Mode2" / "Excel" / EQUIP_XLSX
    if not mode2_equip.is_file():
        raise SystemExit(f"missing {mode2_equip}")
    patch_equip_xlsx(mode2_equip)
    bake_one(ROOT / "Mode2" / "Excel", ROOT / "Mode2" / "Csv", EQUIP_XLSX)


if __name__ == "__main__":
    main()
