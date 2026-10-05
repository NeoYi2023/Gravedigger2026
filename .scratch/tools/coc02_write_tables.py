"""Slice 02: write Mode2 COC sample tables and fill Sandbox COC GameplayConfigId.

Excel keeps the three-row header. CSV is the bake shape (English header + data).
"""
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(r"F:\CursorGame_Git\Gravedigger2026\Gravedigger2026\Assets\ConfigTables")
MODE2_EXCEL = ROOT / "Mode2" / "Excel"
MODE2_CSV = ROOT / "Mode2" / "Csv"

COC_IDS = {
    "SB_N_Coc": "Coc_Normal",
    "SB_H_Coc": "Coc_Hard",
    "SB_L_Coc": "Coc_Hell",
}

TABLES = [
    (
        "COC_玩法配置表_Coc_CocGameplayConfig.xlsx",
        "Coc_CocGameplayConfig.csv",
        ["玩法配置ID", "地图编号", "临时揭示半径", "未探索透明度", "变暗透明度", "卡牌长按秒数", "按住洒兵间隔"],
        [
            "主键；沙盘 CocCombat 行引用",
            "Prefab 逻辑名 Coc_*",
            "地面 XZ 圆半径；缺省 2；≤0 加载失败",
            "缺省 0.9；须在 (0, 1]",
            "缺省 0.7；须在 (0, 1]",
            "缺省 2；≤0 加载失败",
            "缺省 0.1；≤0 加载失败",
        ],
        ["GameplayConfigId", "MapId", "RevealRadius", "UnexploredAlpha", "ExploredAlpha", "CardHoldSeconds", "DeployHoldIntervalSeconds"],
        [
            ["Coc_Normal", "Coc_Lv2_01", "2", "0.9", "0.7", "2", "0.1"],
            ["Coc_Hard", "Coc_Lv2_01", "2", "0.9", "0.7", "2", "0.1"],
            ["Coc_Hell", "Coc_Lv2_01", "2", "0.9", "0.7", "2", "0.1"],
        ],
    ),
    (
        "COC_刷怪配置表_Coc_CocSpawnConfig.xlsx",
        "Coc_CocSpawnConfig.csv",
        ["玩法配置ID", "刷怪点编号", "怪物ID", "数量", "刷怪角色", "同点顺序"],
        [
            "外键 → CocGameplayConfig",
            "与地图 SpawnPoint 匹配",
            "外键 → MonsterConfig",
            "≥1；FinalBoss 必须为 1",
            "Normal / MiniBoss / FinalBoss",
            "同 SpawnPointId 多行时升序",
        ],
        ["GameplayConfigId", "SpawnPointId", "MonsterId", "SpawnCount", "SpawnRole", "SpawnOrder"],
        [
            [gameplay, point, "Monster_01", count, role, "1"]
            for gameplay in ("Coc_Normal", "Coc_Hard", "Coc_Hell")
            for point, count, role in (
                ("SP_01", "3", "Normal"),
                ("SP_Mini", "1", "MiniBoss"),
                ("SP_Final", "1", "FinalBoss"),
            )
        ],
    ),
    (
        "COC_占领点配置表_Coc_CocCapturePointConfig.xlsx",
        "Coc_CocCapturePointConfig.csv",
        ["占领点ID", "玩法配置ID", "永久亮起半径"],
        [
            "主键；与地图 CocCapturePoint 匹配",
            "外键 → CocGameplayConfig",
            "地面 XZ 圆；> 0",
        ],
        ["CapturePointId", "GameplayConfigId", "Radius"],
        [
            ["CP_Normal", "Coc_Normal", "3"],
            ["CP_Hard", "Coc_Hard", "3"],
            ["CP_Hell", "Coc_Hell", "3"],
        ],
    ),
    (
        "COC_占领奖励配置表_Coc_CocCaptureRewardConfig.xlsx",
        "Coc_CocCaptureRewardConfig.csv",
        ["奖励ID", "占领点ID", "目标沙盘节点", "增加次数"],
        [
            "主键",
            "外键 → CocCapturePointConfig",
            "外键 → SandboxNodeConfig.NodeId；须同难度且为挖坟/商店/自动造兵",
            "≥ 1",
        ],
        ["RewardId", "CapturePointId", "TargetNodeId", "AddCount"],
        [
            ["R_N", "CP_Normal", "SB_N_Dig", "1"],
            ["R_H", "CP_Hard", "SB_H_Shop", "1"],
            ["R_L", "CP_Hell", "SB_L_AM", "1"],
        ],
    ),
]


def write_workbook(path: Path, zh, note, en, data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(zh)
    ws.append(note)
    ws.append(en)
    for row in data:
        ws.append(list(row))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_csv(path: Path, en, data):
    lines = [",".join(en)]
    for row in data:
        lines.append(",".join(row))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def patch_sandbox(excel_path: Path, csv_path: Path):
    wb = load_workbook(excel_path)
    ws = wb.active
    header_row = None
    for r in range(1, 4):
        values = [str(c.value).strip() if c.value is not None else "" for c in ws[r]]
        if "GameplayConfigId" in values and "NodeId" in values:
            header_row = r
            headers = values
            break
    if header_row is None:
        raise SystemExit(f"No English header in {excel_path}")

    node_col = headers.index("NodeId") + 1
    config_col = headers.index("GameplayConfigId") + 1
    changed = 0
    for r in range(header_row + 1, ws.max_row + 1):
        node_id = ws.cell(r, node_col).value
        if node_id in COC_IDS:
            ws.cell(r, config_col).value = COC_IDS[node_id]
            changed += 1
    if changed != 3:
        raise SystemExit(f"{excel_path} updated {changed} COC rows, expected 3")
    wb.save(excel_path)

    # Rebuild CSV from the English header downward so bake shape stays aligned.
    en = []
    for c in range(1, ws.max_column + 1):
        text = ws.cell(header_row, c).value
        if text is None or str(text).strip() == "":
            break
        en.append(str(text).strip())
    data = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = []
        empty = True
        for c in range(1, len(en) + 1):
            value = ws.cell(r, c).value
            text = "" if value is None else str(value).strip()
            if text:
                empty = False
            row.append(text)
        if not empty:
            data.append(row)
    write_csv(csv_path, en, data)
    print(f"Patched {excel_path.name} ({changed} COC ids) → {csv_path.name}")


def main():
    for excel_name, csv_name, zh, note, en, data in TABLES:
        excel_path = MODE2_EXCEL / excel_name
        csv_path = MODE2_CSV / csv_name
        write_workbook(excel_path, zh, note, en, data)
        write_csv(csv_path, en, data)
        print(f"Wrote {excel_name} ({len(data)} rows)")

    for excel_dir, csv_dir in (
        (ROOT / "Excel", ROOT / "Csv"),
        (MODE2_EXCEL, MODE2_CSV),
    ):
        patch_sandbox(
            excel_dir / "关卡_沙盘节点表_Level_SandboxNodeConfig.xlsx",
            csv_dir / "Level_SandboxNodeConfig.csv",
        )


if __name__ == "__main__":
    main()
