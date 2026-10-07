# -*- coding: utf-8 -*-
"""D-100 slice 01: ClassConfig observe/score columns + MonsterConfig.TargetValue (SPEC_04 §14.7)."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "Gravedigger2026" / "Assets" / "ConfigTables"
CLASS_XLSX = "制造_职业配置表_Manufacture_ClassConfig.xlsx"
MONSTER_XLSX = "防守_怪物配置表_Defend_MonsterConfig.xlsx"

CLASS_COLS = (
    ("ObserveRange", "观察范围", "仅COC：选敌圆与迷雾半径（世界单位）；缺列/空/≤0→2"),
    ("DistancePriority", "距离优先级", "仅COC：优先最近|优先最远；缺列/空→优先最近；非法失败"),
    ("TargetValueMode", "目标价值模式", "仅COC：由大到小|无视价值；缺列/空→由大到小；非法失败"),
    ("TargetTypeScores", "目标类型分数", "仅COC：类型名;分数|…；类型仅敌方单位|可破坏障碍物；空=无候选"),
)

MONSTER_COLS = (
    ("TargetValue", "目标价值", "仅COC选敌价值分；缺列/空→1；<0失败。障碍价值不在本表"),
)

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bake_config_tables_py import (  # noqa: E402
    build_csv,
    excel_to_csv_base,
    find_header_index,
    read_sheet_rows,
)


def find_en_header_row(ws) -> int:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True), 1):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if any(c in ("ClassId", "MonsterId") for c in cells):
            return i
    raise RuntimeError("no EN header in first 3 rows")


def header_list(ws, header_row: int) -> list[str]:
    row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    names = [str(c).strip() if c is not None else "" for c in row]
    while names and names[-1] == "":
        names.pop()
    return names


def append_column(ws, header_row: int, en: str, zh: str, note: str) -> None:
    names = header_list(ws, header_row)
    if en in names:
        return
    insert_at = len(names) + 1
    ws.insert_cols(insert_at)
    if header_row >= 3:
        ws.cell(1, insert_at, zh)
        ws.cell(2, insert_at, note)
    ws.cell(header_row, insert_at, en)


def col_index(names: list[str], en: str) -> int:
    return names.index(en) + 1


def patch_class(path: Path, fill_demo: bool) -> None:
    wb = load_workbook(path)
    ws = wb.active
    header_row = find_en_header_row(ws)
    for en, zh, note in CLASS_COLS:
        append_column(ws, header_row, en, zh, note)
    names = header_list(ws, header_row)
    id_col = col_index(names, "ClassId")
    cols = {en: col_index(names, en) for en, _, _ in CLASS_COLS}
    for r in range(header_row + 1, ws.max_row + 1):
        class_id = ws.cell(r, id_col).value
        if class_id is None or not str(class_id).strip():
            continue
        if fill_demo:
            ws.cell(r, cols["ObserveRange"], 2)
            ws.cell(r, cols["DistancePriority"], "优先最近")
            ws.cell(r, cols["TargetValueMode"], "由大到小")
            ws.cell(r, cols["TargetTypeScores"], "敌方单位;1")
        else:
            for en, _, _ in CLASS_COLS:
                ws.cell(r, cols[en], "")
    wb.save(path)
    print(f"OK class fill_demo={fill_demo} {path}")


def patch_monster(path: Path, fill_demo: bool) -> None:
    wb = load_workbook(path)
    ws = wb.active
    header_row = find_en_header_row(ws)
    for en, zh, note in MONSTER_COLS:
        append_column(ws, header_row, en, zh, note)
    names = header_list(ws, header_row)
    id_col = col_index(names, "MonsterId")
    value_col = col_index(names, "TargetValue")
    for r in range(header_row + 1, ws.max_row + 1):
        monster_id = ws.cell(r, id_col).value
        if monster_id is None or not str(monster_id).strip():
            continue
        ws.cell(r, value_col, 1 if fill_demo else "")
    wb.save(path)
    print(f"OK monster fill_demo={fill_demo} {path}")


def bake_one(excel_dir: Path, csv_dir: Path, xlsx_name: str) -> None:
    xlsx = excel_dir / xlsx_name
    rows = read_sheet_rows(xlsx)
    header_i = find_header_index(rows)
    text = build_csv(rows[header_i:])
    out = csv_dir / f"{excel_to_csv_base(xlsx.stem)}.csv"
    out.write_text(text, encoding="utf-8", newline="\n")
    header = text.splitlines()[0] if text else ""
    print(f"BAKE {xlsx.name} → {out.name} header={header}")


def main() -> None:
    mode1_excel = ROOT / "Excel"
    mode1_csv = ROOT / "Csv"
    mode2_excel = ROOT / "Mode2" / "Excel"
    mode2_csv = ROOT / "Mode2" / "Csv"

    patch_class(mode1_excel / CLASS_XLSX, fill_demo=False)
    patch_monster(mode1_excel / MONSTER_XLSX, fill_demo=False)
    patch_class(mode2_excel / CLASS_XLSX, fill_demo=True)
    patch_monster(mode2_excel / MONSTER_XLSX, fill_demo=True)

    for excel_dir, csv_dir in ((mode1_excel, mode1_csv), (mode2_excel, mode2_csv)):
        bake_one(excel_dir, csv_dir, CLASS_XLSX)
        bake_one(excel_dir, csv_dir, MONSTER_XLSX)


if __name__ == "__main__":
    main()
