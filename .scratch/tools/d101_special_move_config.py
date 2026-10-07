# -*- coding: utf-8 -*-
"""D-101 slice 01: ClassConfig.SpecialMove column (SPEC_04 §14.7)."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "Gravedigger2026" / "Assets" / "ConfigTables"
CLASS_XLSX = "制造_职业配置表_Manufacture_ClassConfig.xlsx"

COL_EN = "SpecialMove"
COL_ZH = "特殊移动"
COL_NOTE = "0=不具备；1=具备。具备时，地图上支持特殊移动的AirWall对该士兵无效（D-101）"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bake_config_tables_py import (  # noqa: E402
    build_csv,
    excel_to_csv_base,
    find_header_index,
    read_sheet_rows,
)


def find_en_header_row(ws) -> int:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True), 1):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if "ClassId" in cells:
            return i
    raise RuntimeError("no EN header in first 3 rows")


def header_list(ws, header_row: int) -> list[str]:
    row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    names = [str(c).strip() if c is not None else "" for c in row]
    while names and names[-1] == "":
        names.pop()
    return names


def append_column(ws, header_row: int, en: str, zh: str, note: str) -> None:
    names = header_list(ws, header_row)
    if en in names:
        return
    insert_at = len(names) + 1
    ws.insert_cols(insert_at)
    if header_row >= 3:
        ws.cell(1, insert_at, zh)
        ws.cell(2, insert_at, note)
    ws.cell(header_row, insert_at, en)


def col_index(names: list[str], en: str) -> int:
    return names.index(en) + 1


def patch_class(path: Path, rogue_value: int) -> None:
    wb = load_workbook(path)
    ws = wb.active
    header_row = find_en_header_row(ws)
    append_column(ws, header_row, COL_EN, COL_ZH, COL_NOTE)
    names = header_list(ws, header_row)
    id_col = col_index(names, "ClassId")
    flag_col = col_index(names, COL_EN)
    for r in range(header_row + 1, ws.max_row + 1):
        class_id = ws.cell(r, id_col).value
        if class_id is None or not str(class_id).strip():
            continue
        value = rogue_value if str(class_id).strip() == "Class_BaseRogue" else 0
        ws.cell(r, flag_col, value)
    wb.save(path)
    print(f"OK class rogue={rogue_value} {path}")


def bake_one(excel_dir: Path, csv_dir: Path, xlsx_name: str) -> None:
    xlsx = excel_dir / xlsx_name
    rows = read_sheet_rows(xlsx)
    header_i = find_header_index(rows)
    text = build_csv(rows[header_i:])
    out = csv_dir / f"{excel_to_csv_base(xlsx.stem)}.csv"
    out.write_text(text, encoding="utf-8", newline="\n")
    header = text.splitlines()[0] if text else ""
    print(f"BAKE {xlsx.name} → {out.name} header={header}")


def main() -> None:
    mode1_excel = ROOT / "Excel"
    mode1_csv = ROOT / "Csv"
    mode2_excel = ROOT / "Mode2" / "Excel"
    mode2_csv = ROOT / "Mode2" / "Csv"

    patch_class(mode1_excel / CLASS_XLSX, rogue_value=0)
    patch_class(mode2_excel / CLASS_XLSX, rogue_value=1)

    bake_one(mode1_excel, mode1_csv, CLASS_XLSX)
    bake_one(mode2_excel, mode2_csv, CLASS_XLSX)


if __name__ == "__main__":
    main()
