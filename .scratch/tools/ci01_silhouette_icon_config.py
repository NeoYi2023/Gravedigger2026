# -*- coding: utf-8 -*-
"""CI-01: ClassConfig/MonsterConfig SilhouetteIconAssetId + Bake (SPEC_04 §9.9b/§9.19/§14.7)."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "Gravedigger2026" / "Assets" / "ConfigTables"
CLASS_XLSX = "制造_职业配置表_Manufacture_ClassConfig.xlsx"
MONSTER_XLSX = "防守_怪物配置表_Defend_MonsterConfig.xlsx"

COL_EN = "SilhouetteIconAssetId"
COL_ZH = "士兵简画图标"
COL_NOTE = "UI-033 战斗指示器简画；仅文件名；Resources/UI/Icons/{Id}；空=空框"

BASE_CLASS_ICON = {
    "战士": "HPPK_UI_Class_BaseWarrior",
    "射手": "HPPK_UI_Class_BaseArcher",
    "法师": "HPPK_UI_Class_BaseMage",
    "刺客": "HPPK_UI_Class_BaseRogue",
    "盗贼": "HPPK_UI_Class_BaseRogue",  # legacy token
}

MONSTER_ICON = {
    "Monster_01": "HPPK_UI_Monster_01",
    "Monster_02": "HPPK_UI_Monster_02",
    "Monster_03": "HPPK_UI_Monster_03",
    "Monster_04": "HPPK_UI_Monster_04",
}
MONSTER_FALLBACK = "HPPK_UI_Monster_01"

sys.path.insert(0, str(Path(__file__).resolve().parent))
from bake_config_tables_py import (  # noqa: E402
    build_csv,
    excel_to_csv_base,
    find_header_index,
    read_sheet_rows,
)


def find_en_header_row(ws) -> int:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True), 1):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if any(c in ("ClassId", "MonsterId") for c in cells):
            return i
    raise RuntimeError("no EN header in first 3 rows")


def header_list(ws, header_row: int) -> list[str]:
    row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    names = [str(c).strip() if c is not None else "" for c in row]
    while names and names[-1] == "":
        names.pop()
    return names


def append_column(ws, header_row: int, en: str, zh: str, note: str) -> None:
    names = header_list(ws, header_row)
    if en in names:
        return
    insert_at = len(names) + 1
    ws.insert_cols(insert_at)
    if header_row >= 3:
        ws.cell(1, insert_at, zh)
        ws.cell(2, insert_at, note)
    ws.cell(header_row, insert_at, en)


def patch_class(path: Path, fill_demo: bool) -> None:
    wb = load_workbook(path)
    ws = wb.active
    header_row = find_en_header_row(ws)
    append_column(ws, header_row, COL_EN, COL_ZH, COL_NOTE)
    names = header_list(ws, header_row)
    id_col = names.index("ClassId") + 1
    icon_col = names.index(COL_EN) + 1
    base_col = names.index("BaseClass") + 1 if "BaseClass" in names else None
    for r in range(header_row + 1, ws.max_row + 1):
        class_id = ws.cell(r, id_col).value
        if class_id is None or not str(class_id).strip():
            continue
        if not fill_demo:
            ws.cell(r, icon_col, "")
            continue
        base = ""
        if base_col is not None:
            raw = ws.cell(r, base_col).value
            base = str(raw).strip() if raw is not None else ""
        ws.cell(r, icon_col, BASE_CLASS_ICON.get(base, ""))
    wb.save(path)
    print(f"OK class fill_demo={fill_demo} {path}")


def patch_monster(path: Path, fill_demo: bool) -> None:
    wb = load_workbook(path)
    ws = wb.active
    header_row = find_en_header_row(ws)
    append_column(ws, header_row, COL_EN, COL_ZH, COL_NOTE)
    names = header_list(ws, header_row)
    id_col = names.index("MonsterId") + 1
    icon_col = names.index(COL_EN) + 1
    for r in range(header_row + 1, ws.max_row + 1):
        monster_id = ws.cell(r, id_col).value
        if monster_id is None or not str(monster_id).strip():
            continue
        if not fill_demo:
            ws.cell(r, icon_col, "")
            continue
        mid = str(monster_id).strip()
        ws.cell(r, icon_col, MONSTER_ICON.get(mid, MONSTER_FALLBACK))
    wb.save(path)
    print(f"OK monster fill_demo={fill_demo} {path}")


def bake_one(excel_dir: Path, csv_dir: Path, xlsx_name: str) -> None:
    xlsx = excel_dir / xlsx_name
    rows = read_sheet_rows(xlsx)
    header_i = find_header_index(rows)
    text = build_csv(rows[header_i:])
    out = csv_dir / f"{excel_to_csv_base(xlsx.stem)}.csv"
    out.write_text(text, encoding="utf-8", newline="\n")
    header = text.splitlines()[0] if text else ""
    has_col = COL_EN in header
    print(f"BAKE {xlsx.name} → {out.name} has_{COL_EN}={has_col}")


def main() -> None:
    mode1_excel = ROOT / "Excel"
    mode1_csv = ROOT / "Csv"
    mode2_excel = ROOT / "Mode2" / "Excel"
    mode2_csv = ROOT / "Mode2" / "Csv"

    patch_class(mode1_excel / CLASS_XLSX, fill_demo=False)
    patch_monster(mode1_excel / MONSTER_XLSX, fill_demo=False)
    patch_class(mode2_excel / CLASS_XLSX, fill_demo=True)
    patch_monster(mode2_excel / MONSTER_XLSX, fill_demo=True)

    for excel_dir, csv_dir in ((mode1_excel, mode1_csv), (mode2_excel, mode2_csv)):
        bake_one(excel_dir, csv_dir, CLASS_XLSX)
        bake_one(excel_dir, csv_dir, MONSTER_XLSX)


if __name__ == "__main__":
    main()
