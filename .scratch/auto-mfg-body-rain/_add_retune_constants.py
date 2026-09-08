# -*- coding: utf-8 -*-
"""Upsert UI-016 retune AutoMfg* keys into CombatConstantConfig, then Bake CSV."""
from __future__ import annotations

import io
import math
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2] / "Gravedigger2026" / "Assets" / "ConfigTables"
EXCEL_NAME = "通用_常量表_Combat_CombatConstantConfig.xlsx"
CSV_NAME = "Combat_CombatConstantConfig.csv"
EN_COL = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")
MAX_SCAN = 3
_MAX_DECIMAL_PLACES = 10

UPSERT_ROWS = [
    (
        "AutoMfgPileFloorYPx",
        "自动制造尸骸堆地板 Y",
        -300,
        "Body pile floor Y in ref pixels",
        "2D 演出层堆地板相对屏幕中心 Y",
    ),
    (
        "AutoMfgSoldierLandYPx",
        "自动制造士兵落地 Y",
        -410,
        "Revived soldier idle land Y px",
        "复活士兵落地相对屏幕中心 Y",
    ),
    (
        "AutoMfgSoldierMaxEdgePx",
        "自动制造士兵固定边长",
        64,
        "StepB mystery/soldier Image sizeDelta W/H",
        "StepB 问号/士兵 Image sizeDelta 宽高（再 × VisualScale）",
    ),
    (
        "AutoMfgSoldierVisualScale",
        "自动制造士兵视觉缩放",
        8,
        "Revive piece localScale XYZ",
        "复活件 localScale（XYZ）",
    ),
    (
        "AutoMfgSoldierShadowOffsetYPx",
        "自动制造士兵影子相对 Y",
        -40,
        "Shadow local Y vs soldier center",
        "影子相对士兵中心的本地 Y",
    ),
    (
        "AutoMfgBodyAngleMaxDeg",
        "自动制造躯体绝对转角上限",
        270,
        "Clamp body Z angle abs deg; zero omega at limit",
        "躯体 Z 转角夹紧 ±本值；触限清角速度",
    ),
]


def format_numeric_for_csv(number: float) -> str:
    if not math.isfinite(number):
        return str(number)
    if abs(number - round(number)) < 1e-9 and abs(number) < 1e15:
        return str(int(round(number)))
    quant = Decimal(1).scaleb(-_MAX_DECIMAL_PLACES)
    rounded_dec = Decimal(repr(number)).quantize(quant, rounding=ROUND_HALF_UP)
    rounded = float(rounded_dec)
    if abs(rounded - round(rounded)) < 1e-12 and abs(rounded) < 1e15:
        return str(int(round(rounded)))
    text = f"{rounded:.{_MAX_DECIMAL_PLACES}f}".rstrip("0").rstrip(".")
    return text if text not in ("", "-") else "0"


def sanitize_embedded_float_noise(text: str) -> str:
    if not text or "." not in text:
        return text

    def repl(match: re.Match[str]) -> str:
        token = match.group(0)
        frac = token.split(".", 1)[1]
        looks_noisy = len(frac) > 10 or "000000" in frac or "999999" in frac
        if not looks_noisy:
            return token
        try:
            number = float(token)
        except ValueError:
            return token
        return format_numeric_for_csv(number)

    return re.sub(r"-?\d+\.\d+", repl, text)


def cell_to_csv_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return format_numeric_for_csv(value)
    if isinstance(value, Decimal):
        return format_numeric_for_csv(float(value))
    return sanitize_embedded_float_noise(str(value))


def is_english_header(row: list[str]) -> bool:
    saw = False
    for cell in row:
        if cell is None or str(cell).strip() == "":
            continue
        saw = True
        if not EN_COL.match(str(cell).strip()):
            return False
    return saw


def find_header_index(rows: list[list[str]]) -> int:
    for i, row in enumerate(rows[:MAX_SCAN]):
        if is_english_header(row):
            return i
    raise ValueError("no English header in first 3 rows")


def escape_csv_field(value: str) -> str:
    if any(ch in value for ch in [",", '"', "\n", "\r"]):
        return '"' + value.replace('"', '""') + '"'
    return value


def build_csv(rows: list[list[str]]) -> str:
    lines: list[str] = []
    for r_i, row in enumerate(rows):
        if r_i > 0 and all(not (c or "").strip() for c in row):
            continue
        lines.append(",".join(escape_csv_field(c or "") for c in row))
    return "\n".join(lines) + ("\n" if lines else "")


def upsert_rows(xlsx_path: Path) -> None:
    wb = load_workbook(xlsx_path)
    ws = wb.active
    header1 = [ws.cell(1, c).value for c in range(1, 6)]
    header2 = [ws.cell(2, c).value for c in range(1, 6)]
    header3 = [ws.cell(3, c).value for c in range(1, 6)]
    if header3 != ["ConstantKey", "ConstantKeyZh", "Value", "Comment", "CommentZh"]:
        raise SystemExit(f"unexpected header row 3 in {xlsx_path}: {header3}")
    if not header1[0] or not header2[0]:
        raise SystemExit(f"missing doc rows 1-2 in {xlsx_path}")

    key_to_row = {}
    for r in range(4, ws.max_row + 1):
        key = ws.cell(r, 1).value
        if key:
            key_to_row[str(key)] = r

    updated = 0
    for row in UPSERT_ROWS:
        key = row[0]
        if key in key_to_row:
            r = key_to_row[key]
            for c, val in enumerate(row, start=1):
                ws.cell(r, c).value = val
            updated += 1

    missing = [row for row in UPSERT_ROWS if row[0] not in key_to_row]
    if missing:
        insert_at = None
        last_auto = None
        for r in range(1, ws.max_row + 1):
            key = ws.cell(r, 1).value
            if key and str(key).startswith("AutoMfg"):
                last_auto = r
        insert_at = (last_auto + 1) if last_auto is not None else ws.max_row + 1
        ws.insert_rows(insert_at, amount=len(missing))
        for i, row in enumerate(missing):
            for c, val in enumerate(row, start=1):
                ws.cell(insert_at + i, c).value = val
        print(f"INSERTED {len(missing)} rows at Excel row {insert_at}: {xlsx_path}")

    if updated:
        print(f"UPDATED {updated} rows: {xlsx_path}")

    if [ws.cell(1, c).value for c in range(1, 6)] != header1:
        raise SystemExit(f"row 1 mutated in {xlsx_path}")
    if [ws.cell(2, c).value for c in range(1, 6)] != header2:
        raise SystemExit(f"row 2 mutated in {xlsx_path}")

    wb.save(xlsx_path)
    wb.close()


def bake(excel_dir: Path, csv_dir: Path) -> None:
    xlsx = excel_dir / EXCEL_NAME
    with open(xlsx, "rb") as raw:
        data = raw.read()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        cells = [cell_to_csv_text(v) for v in row]
        rows.append(cells)
    wb.close()
    if not rows:
        raise SystemExit(f"empty sheet {xlsx}")
    max_cols = max(len(r) for r in rows)
    for r in rows:
        if len(r) < max_cols:
            r.extend([""] * (max_cols - len(r)))
    header_i = find_header_index(rows)
    if header_i < 2:
        raise SystemExit(f"expected 3-row header, got English at row {header_i} in {xlsx}")
    text = build_csv(rows[header_i:])
    out = csv_dir / CSV_NAME
    csv_dir.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8", newline="\n")
    keys = [ln.split(",", 1)[0] for ln in text.splitlines()[1:]]
    missing = [k[0] for k in UPSERT_ROWS if k[0] not in keys]
    if missing:
        raise SystemExit(f"Bake missing keys {missing} in {out}")
    print(f"BAKED {xlsx.name} → {out}")


def main() -> int:
    targets = [
        (ROOT / "Excel" / EXCEL_NAME, ROOT / "Excel", ROOT / "Csv"),
        (ROOT / "Mode2" / "Excel" / EXCEL_NAME, ROOT / "Mode2" / "Excel", ROOT / "Mode2" / "Csv"),
    ]
    for xlsx, excel_dir, csv_dir in targets:
        if not xlsx.is_file():
            raise SystemExit(f"missing {xlsx}")
        upsert_rows(xlsx)
        bake(excel_dir, csv_dir)
    return 0


if __name__ == "__main__":
    sys.exit(main())
